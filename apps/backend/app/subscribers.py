"""Private registry API and authenticated bot account/reading operations."""
import csv
import io
import re
import zipfile
from datetime import date, timedelta
from decimal import Decimal
from typing import Annotated, Literal
from uuid import UUID

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, Response
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import Field, ValidationError, field_validator, model_validator
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from starlette.concurrency import run_in_threadpool

from app.auth import bot_auth, rate_limit, super_admin
from app.database import get_db, utcnow
from app.models import Admin, Application, ApplicationFile, AuditLog, FileType
from app.repositories import application_dict
from app.schemas import StrictModel, local_today
from app.services import audit
from app.storage import storage_path
from app.subscriber_models import MeterReading, SubscriberAccount
from app.vision import looks_like_gas_meter

router = APIRouter(prefix="/api/subscribers", tags=["subscribers"])
internal = APIRouter(prefix="/api/internal/subscribers", dependencies=[Depends(bot_auth)])
DB = Annotated[AsyncSession, Depends(get_db)]
Super = Annotated[Admin, Depends(super_admin)]


def normalize_phone(value: str | None) -> str | None:
    if not value:
        return None
    if not re.fullmatch(r"[+0-9()\s-]+", value):
        raise ValueError("Телефонды +7XXXXXXXXXX форматымен енгізіңіз")
    digits = re.sub(r"[^0-9]", "", value)
    if len(digits) == 10:
        digits = "7" + digits
    if len(digits) == 11 and digits.startswith("8"):
        digits = "7" + digits[1:]
    if not re.fullmatch(r"7[0-9]{10}", digits):
        raise ValueError("Телефонды +7XXXXXXXXXX форматымен енгізіңіз")
    return "+" + digits


class AccountInput(StrictModel):
    account_number: str = Field(pattern=r"^[0-9]{6,20}$")
    meter_number: str = Field(min_length=3, max_length=40, pattern=r"^[A-Za-z0-9-]+$")
    full_name: str = Field(min_length=2, max_length=200)
    address: str = Field(min_length=3, max_length=500)
    phone: str | None = Field(default=None, max_length=32)

    @field_validator("phone")
    @classmethod
    def phone_format(cls, value):
        return normalize_phone(value)

    @field_validator("meter_number")
    @classmethod
    def meter_format(cls, value):
        return value.upper()


class AccountUpdate(AccountInput):
    version: int = Field(ge=1)
    active: bool


MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000
IMPORT_HEADERS = ["Дербес шот", "Есептегіш нөмірі", "Аты-жөні", "Мекенжай", "Телефон"]
IMPORT_LABELS = {
    "дербес шот": "account_number", "лицевой счет": "account_number", "лицевой счёт": "account_number",
    "есептегіш нөмірі": "meter_number", "есептегіш": "meter_number", "номер счетчика": "meter_number", "счетчик": "meter_number",
    "аты-жөні": "full_name", "фио": "full_name",
    "мекенжай": "address", "адрес": "address",
    "телефон": "phone", "phone": "phone",
}


def import_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Абоненттер"
    sheet.append(IMPORT_HEADERS)
    sheet.append(["0012345", "M-000123", "Тестов Т.Т.", "Атакент/АСЫКАТА/Мысал/1", "+77010000000"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def rows_from_matrix(matrix: list[list[str]]) -> list[dict]:
    if not matrix:
        return []
    header = [str(cell).strip().lower() for cell in matrix[0]]
    index = {IMPORT_LABELS[label]: i for i, label in enumerate(header) if label in IMPORT_LABELS}
    missing = [f for f in ("account_number", "meter_number", "full_name", "address") if f not in index]
    if missing:
        raise HTTPException(422, "Файлдың бірінші жолында тақырыптар болуы керек: "
                            "Дербес шот, Есептегіш нөмірі, Аты-жөні, Мекенжай (Телефон міндетті емес). "
                            "Үлгіні жүктеп алып, соны толтырыңыз.")

    def cell(raw_row, key):
        i = index.get(key)
        return str(raw_row[i]).strip() if i is not None and i < len(raw_row) and raw_row[i] is not None else ""

    rows = []
    for line, raw_row in enumerate(matrix[1:], start=2):
        if not any(str(c or "").strip() for c in raw_row):
            continue
        rows.append({"line": line, "account_number": cell(raw_row, "account_number"),
                     "meter_number": cell(raw_row, "meter_number"), "full_name": cell(raw_row, "full_name"),
                     "address": cell(raw_row, "address"), "phone": cell(raw_row, "phone") or None})
    return rows


def parse_csv_rows(raw: bytes) -> list[dict]:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise HTTPException(422, "Файлдың кодировкасын анықтау мүмкін болмады. UTF-8 немесе Windows-1251 қолданыңыз.")
    return rows_from_matrix(list(csv.reader(io.StringIO(text))))


def parse_xlsx_rows(raw: bytes) -> list[dict]:
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except (InvalidFileException, KeyError, zipfile.BadZipFile) as exc:
        raise HTTPException(422, "Файл бүлінген немесе .xlsx емес.") from exc
    matrix = []
    for row in workbook.worksheets[0].iter_rows():
        cells = []
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                raise HTTPException(422, f"{c.coordinate} ұяшығында формула бар. Формулалар қабылданбайды.")
            cells.append(c.value)
        matrix.append(cells)
    return rows_from_matrix(matrix)


def parse_import_rows(filename: str, raw: bytes) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return parse_csv_rows(raw)
    if name.endswith(".xlsx"):
        return parse_xlsx_rows(raw)
    raise HTTPException(415, "Тек .xlsx немесе .csv файл қабылданады")


class BindingUpdate(StrictModel):
    telegram_user_id: int | None = Field(default=None, gt=0, le=2**52)
    version: int = Field(ge=1)
    reason: str = Field(min_length=10, max_length=300)


class AccountLookup(StrictModel):
    telegram_user_id: int = Field(gt=0, le=2**52)
    identifier: str = Field(min_length=3, max_length=40, pattern=r"^[A-Za-z0-9-]+$")
    kind: Literal["account", "meter"] = "account"


class ContactProof(AccountLookup):
    contact_user_id: int = Field(gt=0, le=2**52)
    phone: str = Field(max_length=32)

    @field_validator("phone")
    @classmethod
    def phone_format(cls, value):
        return normalize_phone(value)


class ReadingInput(StrictModel):
    account_id: int = Field(gt=0)
    telegram_user_id: int = Field(gt=0, le=2**52)
    idempotency_key: UUID
    photo_id: UUID
    period: date

    @model_validator(mode="after")
    def current_period(self):
        if self.period != local_today().replace(day=1):
            raise ValueError("Көрсеткіш ағымдағы ай үшін қабылданады")
        return self


class ReadingReview(StrictModel):
    status: Literal["ACCEPTED", "REJECTED"]
    note: str = Field(min_length=3, max_length=500)
    value: Decimal | None = Field(default=None, ge=0, le=99999999999, max_digits=14, decimal_places=3)

    @model_validator(mode="after")
    def value_required_on_accept(self):
        if self.status == "ACCEPTED" and self.value is None:
            raise ValueError("Қабылдау үшін көрсеткіш мәнін енгізіңіз")
        return self


def field_value(row, key):
    value = getattr(row, key)
    return str(value) if isinstance(value, Decimal) else value


def account_dict(row):
    return {key: field_value(row, key) for key in (
        "id", "account_number", "meter_number", "full_name", "address", "phone", "active",
        "telegram_user_id", "binding_note", "version", "last_reading", "last_period",
    )}


def resident_dict(row):
    # Personal-account-number mode: anyone who enters a valid account number sees this.
    return {key: field_value(row, key) for key in (
        "id", "account_number", "meter_number", "full_name", "address", "last_reading", "last_period",
    )}


def reading_dict(row):
    return {key: field_value(row, key) for key in (
        "id", "account_id", "period", "value", "consumption", "status", "review_note", "created_at",
    )}


async def lookup(db, data, lock=False):
    column = SubscriberAccount.account_number if data.kind == "account" else SubscriberAccount.meter_number
    query = select(SubscriberAccount).where(column == data.identifier.upper(), SubscriberAccount.active.is_(True))
    return await db.scalar(query.with_for_update() if lock else query)


async def owned_account(db, account_id, telegram_user_id):
    row = await db.scalar(select(SubscriberAccount).where(
        SubscriberAccount.id == account_id,
        SubscriberAccount.telegram_user_id == telegram_user_id,
        SubscriberAccount.active.is_(True),
    ).with_for_update())
    if row is None:
        raise HTTPException(403, "Шотқа қолжетімділікті қайта растаңыз: /start")
    return row


@router.get("/stats")
async def registry_stats(db: DB, admin: Super):
    period = local_today().replace(day=1)
    bound = SubscriberAccount.active.is_(True) & SubscriberAccount.telegram_user_id.isnot(None)
    total = await db.scalar(select(func.count()).select_from(SubscriberAccount).where(bound))
    accepted = await db.scalar(select(func.count()).select_from(SubscriberAccount)
        .where(bound, SubscriberAccount.last_period == period))
    pending = await db.scalar(select(func.count(func.distinct(MeterReading.account_id)))
        .select_from(MeterReading).join(SubscriberAccount)
        .where(bound, MeterReading.period == period, MeterReading.status == "PENDING"))
    return {"period": period.isoformat(), "total": total, "accepted": accepted, "pending": pending,
            "missing": total - accepted - pending}


def safe_cell(value):
    value = str(value if value is not None else "")
    return "'" + value if value.startswith(("=", "+", "-", "@", "\t", "\r")) else value


def make_accounts_export(rows, fmt):
    header = ["Дербес шот", "Есептегіш нөмірі", "Аты-жөні", "Мекенжай", "Телефон", "Күйі", "Telegram",
              "Соңғы қабылданған көрсеткіш", "Соңғы кезең"]
    data = [[
        safe_cell(row.account_number), safe_cell(row.meter_number), safe_cell(row.full_name), safe_cell(row.address),
        safe_cell(row.phone or ""), "Белсенді" if row.active else "Бұғатталған",
        "Расталған" if row.telegram_user_id else "Расталмаған",
        str(row.last_reading) if row.last_reading is not None else "",
        row.last_period.isoformat() if row.last_period else "",
    ] for row in rows]
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(header)
        writer.writerows(data)
        return buffer.getvalue().encode("utf-8-sig")
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Абоненттер")
    sheet.append(header)
    for row in data:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@router.get("/export")
async def export_accounts(db: DB, admin: Super, format: Literal["csv", "xlsx"] = "csv"):
    await rate_limit(f"subscriber-export:{admin.id}", 10, 300)
    rows = (await db.scalars(select(SubscriberAccount).order_by(SubscriberAccount.account_number))).all()
    result = await run_in_threadpool(make_accounts_export, rows, format)
    audit(db, admin.id, "subscriber_export", format)
    await db.commit()
    media = ("text/csv; charset=utf-8" if format == "csv"
             else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return Response(result, media_type=media,
                    headers={"Content-Disposition": f'attachment; filename="abonentter.{format}"'})


@router.get("")
async def accounts(db: DB, admin: Super, q: str = Query("", max_length=100), page: int = Query(1, ge=1)):
    query = select(SubscriberAccount)
    if q:
        # Treat user search characters literally.
        query = query.where(or_(*[column.contains(q.upper() if column.key == "meter_number" else q, autoescape=True)
            for column in (SubscriberAccount.account_number, SubscriberAccount.meter_number, SubscriberAccount.full_name)]))
    total = await db.scalar(select(func.count()).select_from(query.subquery()))
    rows = (await db.scalars(query.order_by(SubscriberAccount.id.desc()).offset((page - 1) * 25).limit(25))).all()
    audit(db, admin.id, "subscriber_list", page)
    await db.commit()
    return {"items": [account_dict(row) for row in rows], "total": total, "page": page}


async def commit_account(db):
    try:
        await db.commit()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(409, "Дербес шот немесе есептегіш нөмірі тіркелген") from exc


@router.get("/import/template")
async def import_template_file(admin: Super):
    return Response(
        content=import_template(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=abonentter_ulgisi.xlsx"},
    )


@router.post("/import")
async def import_accounts(db: DB, admin: Super, file: UploadFile = File()):
    raw = await file.read(MAX_IMPORT_BYTES + 1)
    if not raw or len(raw) > MAX_IMPORT_BYTES:
        raise HTTPException(413, "Файл көлемі 5 МБ-тан аспауы керек")
    rows = parse_import_rows(file.filename or "", raw)
    if len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(422, f"Файлда {MAX_IMPORT_ROWS}-ден көп жол бар")
    if not rows:
        raise HTTPException(422, "Файлда деректер жоқ")
    created, skipped, errors = [], [], []
    seen_accounts, seen_meters = set(), set()
    for row in rows:
        try:
            data = AccountInput(account_number=row["account_number"], meter_number=row["meter_number"],
                                full_name=row["full_name"], address=row["address"], phone=row["phone"])
        except ValidationError as exc:
            errors.append({"line": row["line"], "reason": "; ".join(e["msg"] for e in exc.errors())})
            continue
        if data.account_number in seen_accounts or data.meter_number in seen_meters:
            errors.append({"line": row["line"], "reason": "Файл ішінде қайталанған дербес шот немесе есептегіш"})
            continue
        seen_accounts.add(data.account_number)
        seen_meters.add(data.meter_number)
        try:
            async with db.begin_nested():
                db.add(SubscriberAccount(**data.model_dump()))
                await db.flush()
        except IntegrityError:
            skipped.append({"line": row["line"], "reason": "Дербес шот немесе есептегіш нөмірі бұрыннан тіркелген"})
            continue
        created.append(row["line"])
    audit(db, admin.id, "subscriber_import", len(created))
    await db.commit()
    return {"created": len(created), "skipped": skipped, "errors": errors, "total": len(rows)}


@router.post("", status_code=201)
async def create_account(data: AccountInput, db: DB, admin: Super):
    row = SubscriberAccount(**data.model_dump())
    db.add(row)
    try:
        await db.flush()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(409, "Дербес шот немесе есептегіш нөмірі тіркелген") from exc
    audit(db, admin.id, "subscriber_create", row.id)
    await commit_account(db)
    return account_dict(row)


@router.put("/{account_id}")
async def update_account(account_id: int, data: AccountUpdate, db: DB, admin: Super):
    row = await db.scalar(select(SubscriberAccount).where(SubscriberAccount.id == account_id).with_for_update())
    if row is None:
        raise HTTPException(404, "Абонент табылмады")
    if row.version != data.version:
        raise HTTPException(409, "Деректер өзгерген. Бетті жаңартыңыз.")
    if row.account_number != data.account_number or row.meter_number != data.meter_number:
        raise HTTPException(422, "Тіркелген шот пен есептегіш нөмірін өзгертуге болмайды")
    # Any identity/contact edit invalidates the previous verification.
    if any(getattr(row, key) != getattr(data, key) for key in ("full_name", "address", "phone", "active")):
        row.telegram_user_id = None
        row.binding_note = None
    for key, value in data.model_dump(exclude={"version"}).items():
        setattr(row, key, value)
    row.version += 1
    audit(db, admin.id, "subscriber_update", row.id)
    await commit_account(db)
    return account_dict(row)


@router.post("/{account_id}/binding")
async def manual_binding(account_id: int, data: BindingUpdate, db: DB, admin: Super):
    row = await db.scalar(select(SubscriberAccount).where(SubscriberAccount.id == account_id).with_for_update())
    if row is None:
        raise HTTPException(404, "Абонент табылмады")
    if row.version != data.version:
        raise HTTPException(409, "Деректер өзгерген. Бетті жаңартыңыз.")
    if not row.active and data.telegram_user_id is not None:
        raise HTTPException(422, "Алдымен абонентті белсендіріңіз")
    row.telegram_user_id = data.telegram_user_id
    row.binding_note = data.reason
    row.version += 1
    # The current verification reason is available to administrators only.
    audit(db, admin.id, "subscriber_bind" if data.telegram_user_id else "subscriber_unbind", row.id)
    await db.commit()
    return account_dict(row)


HISTORY_LABELS = {
    "subscriber_create": "Абонент тіркелді", "subscriber_update": "Деректер өзгертілді",
    "subscriber_bind": "Telegram қолмен байланыстырылды", "subscriber_unbind": "Telegram байланысы жойылды",
    "subscriber_number_only_bind": "Дербес шот нөмірі арқылы автобайланыс",
    "subscriber_contact_verified": "Телефон арқылы расталды",
}


@router.get("/{account_id}/profile")
async def account_profile(account_id: int, db: DB, admin: Super):
    account = await db.scalar(select(SubscriberAccount).where(SubscriberAccount.id == account_id))
    if account is None:
        raise HTTPException(404, "Абонент табылмады")
    readings = (await db.scalars(select(MeterReading).where(MeterReading.account_id == account_id)
        .order_by(MeterReading.created_at.desc()))).all()
    applications = (await db.scalars(select(Application).where(Application.personal_account == account.account_number)
        .order_by(Application.created_at.desc()))).all()
    history_rows = (await db.execute(select(AuditLog, Admin.name).outerjoin(Admin, AuditLog.admin_id == Admin.id)
        .where(AuditLog.target == str(account_id), AuditLog.action.in_(HISTORY_LABELS))
        .order_by(AuditLog.created_at.desc()))).all()
    audit(db, admin.id, "subscriber_profile_view", account_id)
    await db.commit()
    return {
        "account": account_dict(account),
        "readings": [reading_dict(r) for r in readings],
        "applications": [await application_dict(db, a) for a in applications],
        "history": [{"action": HISTORY_LABELS[log.action], "admin_name": name or "Bot",
                     "created_at": log.created_at.isoformat()} for log, name in history_rows],
    }


@internal.post("/lookup")
async def resident_lookup(data: AccountLookup, db: DB):
    # No phone/ID check in this mode: a correct account number is treated as sufficient proof,
    # and the first Telegram account to look it up is bound to it.
    await rate_limit(f"subscriber-lookup:{data.telegram_user_id}", 15, 900)
    row = await lookup(db, data, lock=True)
    if row is None or row.telegram_user_id not in (None, data.telegram_user_id):
        return {"verified": False}
    if row.telegram_user_id is None:
        row.telegram_user_id = data.telegram_user_id
        row.binding_note = "Auto-confirmed by personal account number only (no phone/ID check)"
        row.version += 1
        audit(db, None, "subscriber_number_only_bind", row.id)
        await db.commit()
    return {"verified": True, "account": resident_dict(row)}


@internal.post("/verify-contact")
async def verify_contact(data: ContactProof, db: DB):
    await rate_limit(f"subscriber-proof:{data.telegram_user_id}", 5, 900)
    row = await lookup(db, data, lock=True)
    if (data.contact_user_id != data.telegram_user_id or row is None or not row.phone
            or row.phone != data.phone or row.telegram_user_id not in (None, data.telegram_user_id)):
        return {"verified": False}
    row.telegram_user_id = data.telegram_user_id
    row.binding_note = "Telegram contact matches registered phone"
    row.version += 1
    audit(db, None, "subscriber_contact_verified", row.id)
    await db.commit()
    return {"verified": True, "account": resident_dict(row)}


@internal.post("/readings", status_code=201)
async def submit_reading(data: ReadingInput, db: DB):
    await rate_limit(f"meter-reading:{data.telegram_user_id}", 20, 3600)
    account = await owned_account(db, data.account_id, data.telegram_user_id)
    existing = await db.scalar(select(MeterReading).where(
        MeterReading.telegram_user_id == data.telegram_user_id,
        MeterReading.idempotency_key == str(data.idempotency_key),
    ))
    if existing:
        if (existing.account_id, existing.period, existing.photo_id) != (data.account_id, data.period, str(data.photo_id)):
            raise HTTPException(409, "Бұл сұраныс басқа есеп үшін қолданылған")
        return reading_dict(existing)
    if account.last_period and data.period <= account.last_period:
        raise HTTPException(409, "Осы айдың көрсеткіші қабылданған")
    pending = await db.scalar(select(MeterReading.id).where(
        MeterReading.account_id == account.id, MeterReading.period == data.period,
        MeterReading.status == "PENDING",
    ))
    if pending:
        raise HTTPException(409, "Осы айдың есебі тексерілуде. Қайта жіберудің қажеті жоқ.")
    photo = await db.scalar(select(ApplicationFile).where(ApplicationFile.id == str(data.photo_id)).with_for_update())
    if photo is None or photo.owner_telegram_id != data.telegram_user_id or photo.application_id is not None \
            or photo.file_type != FileType.METER_READING_PHOTO:
        raise HTTPException(422, "Фото табылмады немесе рұқсат жоқ. Қайта жіберіңіз.")
    created = photo.created_at if photo.created_at.tzinfo else photo.created_at.replace(tzinfo=utcnow().tzinfo)
    if utcnow() - created > timedelta(hours=24):
        raise HTTPException(422, "Фото мерзімі өткен. Қайта жіберіңіз.")
    already_used = await db.scalar(select(MeterReading.id).where(MeterReading.photo_id == str(data.photo_id)))
    if already_used:
        raise HTTPException(422, "Бұл фото бұрын пайдаланылған. Қайта түсіріп жіберіңіз.")
    if photo.content_hash:
        reused = await db.scalar(select(MeterReading.id).join(ApplicationFile, MeterReading.photo_id == ApplicationFile.id)
            .where(MeterReading.account_id == data.account_id, ApplicationFile.content_hash == photo.content_hash))
        if reused:
            raise HTTPException(422, "Бұл фото осы шот бойынша ертерек жіберілген. "
                                     "Есептегіштің жаңа, ағымдағы фотосын түсіріп жіберіңіз.")
    path = storage_path(photo.storage_url)
    if path.is_file():
        raw = await run_in_threadpool(path.read_bytes)
        if await looks_like_gas_meter(raw) is False:
            raise HTTPException(422, "Фотода есептегіш көрінбейді. Есептегіштің нақты фотосын жіберіңіз.")
    row = MeterReading(account_id=data.account_id, telegram_user_id=data.telegram_user_id,
                        idempotency_key=str(data.idempotency_key), period=data.period, photo_id=str(data.photo_id))
    db.add(row)
    try:
        await db.flush()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(409, "Бұл сұраныс бұрын қолданылған. /start арқылы қайта кіріңіз.") from exc
    audit(db, None, "meter_reading_submit", row.id)
    await db.commit()
    return reading_dict(row)


@internal.get("/{telegram_user_id}/readings")
async def own_readings(telegram_user_id: int, db: DB):
    await rate_limit(f"meter-history:{telegram_user_id}", 30, 60)
    rows = (await db.execute(select(MeterReading, SubscriberAccount).join(SubscriberAccount).where(
        MeterReading.telegram_user_id == telegram_user_id,
        SubscriberAccount.telegram_user_id == telegram_user_id,
        SubscriberAccount.active.is_(True),
    ).order_by(MeterReading.created_at.desc()).limit(10))).all()
    return {"items": [{**reading_dict(reading), "account_number": account.account_number} for reading, account in rows]}


@router.get("/readings/list")
async def reading_list(db: DB, admin: Super, page: int = Query(1, ge=1)):
    rows = (await db.execute(select(MeterReading, SubscriberAccount).join(SubscriberAccount)
        .order_by(MeterReading.created_at.desc()).offset((page - 1) * 25).limit(25))).all()
    total = await db.scalar(select(func.count()).select_from(MeterReading))
    audit(db, admin.id, "meter_reading_list", page)
    await db.commit()
    return {"items": [{**reading_dict(reading), "account_number": account.account_number,
                       "meter_number": account.meter_number, "full_name": account.full_name,
                       "last_reading": field_value(account, "last_reading")} for reading, account in rows], "total": total}


@router.get("/readings/{reading_id}/photo")
async def reading_photo(reading_id: UUID, db: DB, admin: Super):
    reading = await db.scalar(select(MeterReading).where(MeterReading.id == str(reading_id)))
    if reading is None:
        raise HTTPException(404, "Есеп табылмады")
    file = await db.get(ApplicationFile, reading.photo_id)
    path = storage_path(file.storage_url) if file else None
    if path is None or not path.is_file():
        raise HTTPException(404, "Фото табылмады")
    return FileResponse(
        path,
        media_type="image/jpeg",
        headers={"Cache-Control": "private, no-store", "X-Content-Type-Options": "nosniff"},
    )


@router.post("/readings/{reading_id}/review")
async def review_reading(reading_id: UUID, data: ReadingReview, db: DB, admin: Super):
    account_id = await db.scalar(select(MeterReading.account_id).where(MeterReading.id == str(reading_id)))
    if account_id is None:
        raise HTTPException(404, "Есеп табылмады")
    account = await db.scalar(select(SubscriberAccount).where(SubscriberAccount.id == account_id).with_for_update())
    row = await db.scalar(select(MeterReading).where(MeterReading.id == str(reading_id)).with_for_update())
    if row.status != "PENDING":
        raise HTTPException(409, "Есеп тексеріліп қойған")
    if data.status == "ACCEPTED":
        if not account.active:
            raise HTTPException(409, "Абонент бұғатталған")
        if account.last_period and row.period <= account.last_period:
            raise HTTPException(409, "Бұл кезең қабылданған немесе ескірген")
        if account.last_reading is not None and data.value < account.last_reading:
            raise HTTPException(422, "Көрсеткіш төмендеген. Фотоны қайта тексеріп, есепті қабылдамаңыз.")
        # Consumption is unknown for the initial baseline, never assumed to be zero.
        row.value = data.value
        row.consumption = None if account.last_reading is None else data.value - account.last_reading
        account.last_reading, account.last_period = data.value, row.period
        account.version += 1
    row.status, row.review_note, row.reviewed_by = data.status, data.note, admin.id
    audit(db, admin.id, "meter_reading_review", row.id)
    await db.commit()
    return reading_dict(row)
