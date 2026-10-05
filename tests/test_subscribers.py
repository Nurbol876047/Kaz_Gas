import uuid
from decimal import Decimal

import pytest
from sqlalchemy import func, select

from app.config import get_settings
from app.models import ApplicationFile, FileType, Role
from app.schemas import local_today
from app.subscriber_models import MeterReading, SubscriberAccount


def bot_headers():
    return {"X-Bot-Key": get_settings().bot_api_key}


async def register(client, **changes):
    data = {"account_number": "00123456", "meter_number": "M-000123", "full_name": "Test subscriber",
            "address": "Synthetic test street 1", "phone": "+77010000001", **changes}
    response = await client.post("/api/subscribers", json=data)
    assert response.status_code == 201, response.text
    return response.json()


async def upload_reading_photo(db, telegram_user_id, content_hash=None):
    photo = ApplicationFile(owner_telegram_id=telegram_user_id, telegram_file_id="tg-" + str(uuid.uuid4()),
                            file_type=FileType.METER_READING_PHOTO, storage_url=f"{uuid.uuid4()}.jpg",
                            content_hash=content_hash)
    db.add(photo)
    await db.commit()
    return str(photo.id)


async def test_registry_requires_admin_and_preserves_identifiers(client, logged_in, admin, db):
    row = await register(logged_in)
    assert row["account_number"] == "00123456" and row["meter_number"] == "M-000123"
    duplicate = await logged_in.post("/api/subscribers", json={key: row[key] for key in
        ("account_number", "meter_number", "full_name", "address", "phone")})
    assert duplicate.status_code == 409
    admin.role = Role.OPERATOR
    await db.commit()
    assert (await client.get("/api/subscribers")).status_code == 403
    assert (await client.get("/api/subscribers/readings/list")).status_code == 403


async def test_lookup_by_account_number_binds_first_asker_only(logged_in):
    row = await register(logged_in)
    unknown = await logged_in.post("/api/internal/subscribers/lookup", headers=bot_headers(), json={
        "telegram_user_id": 777, "identifier": "99999999", "kind": "account"})
    assert unknown.status_code == 200 and unknown.json() == {"verified": False}
    found = await logged_in.post("/api/internal/subscribers/lookup", headers=bot_headers(), json={
        "telegram_user_id": 777, "identifier": row["account_number"], "kind": "account"})
    assert found.status_code == 200
    body = found.json()["account"]
    assert found.json()["verified"] is True and body["full_name"] == "Test subscriber"
    # A different Telegram user cannot take over the account once it is bound.
    other = await logged_in.post("/api/internal/subscribers/lookup", headers=bot_headers(), json={
        "telegram_user_id": 888, "identifier": row["account_number"], "kind": "account"})
    assert other.json() == {"verified": False}
    assert (await logged_in.post("/api/internal/subscribers/lookup", json={
        "telegram_user_id": 777, "identifier": "00123456"})).status_code == 401


async def test_own_contact_binding_and_revoke(logged_in):
    row = await register(logged_in)
    body = {"telegram_user_id": 777, "contact_user_id": 777, "identifier": "M-000123", "kind": "meter",
            "phone": "87010000001"}
    for changes in [{"contact_user_id": 888}, {"phone": "+77010000002"}]:
        response = await logged_in.post("/api/internal/subscribers/verify-contact", headers=bot_headers(),
                                        json={**body, **changes})
        assert response.json() == {"verified": False}
    proof = await logged_in.post("/api/internal/subscribers/verify-contact", headers=bot_headers(), json=body)
    assert proof.json()["account"]["full_name"] == "Test subscriber"
    assert "phone" not in proof.json()["account"]
    # Another Telegram account cannot take over even with the same claimed phone.
    other = await logged_in.post("/api/internal/subscribers/verify-contact", headers=bot_headers(),
                                 json={**body, "telegram_user_id": 888, "contact_user_id": 888})
    assert other.json() == {"verified": False}
    updated = await logged_in.put(f"/api/subscribers/{row['id']}", json={
        **{key: row[key] for key in ("account_number", "meter_number", "full_name", "address", "phone")},
        "version": row["version"] + 1, "active": False})
    assert updated.status_code == 200 and updated.json()["telegram_user_id"] is None
    assert (await logged_in.post("/api/internal/subscribers/lookup", headers=bot_headers(), json={
        "telegram_user_id": 777, "identifier": "00123456"})).json() == {"verified": False}


def import_xlsx_bytes(rows):
    import io
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Дербес шот", "Есептегіш нөмірі", "Аты-жөні", "Мекенжай", "Телефон"])
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def test_import_template_download(logged_in):
    response = await logged_in.get("/api/subscribers/import/template")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")


async def test_import_requires_admin(client):
    response = await client.post("/api/subscribers/import", files={"file": ("x.csv", b"data", "text/csv")})
    assert response.status_code == 401


async def test_import_rejects_wrong_extension(logged_in):
    response = await logged_in.post("/api/subscribers/import",
                                    files={"file": ("x.txt", b"data", "text/plain")})
    assert response.status_code == 415


async def test_import_rejects_xlsx_formula(logged_in):
    content = import_xlsx_bytes([["0111111", "M-IMP-1", "Тестов Т.", "Мекенжай 1", ""]])
    import io
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(content))
    workbook.active["A2"] = "=1+1"
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = await logged_in.post("/api/subscribers/import",
        files={"file": ("import.xlsx", buffer.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 422
    assert "формула" in response.text.lower()


async def test_import_requires_known_headers(logged_in):
    content = "Foo,Bar\n123,456\n".encode()
    response = await logged_in.post("/api/subscribers/import", files={"file": ("x.csv", content, "text/csv")})
    assert response.status_code == 422
    assert "тақырыптар" in response.text.lower()


async def test_import_creates_skips_duplicates_and_reports_errors(logged_in, db):
    await register(logged_in, account_number="0777777", meter_number="M-EXIST")
    content = import_xlsx_bytes([
        ["0111111", "M-IMP-1", "Бірінші Т.", "Мекенжай 1", "+77010000010"],
        ["0111111", "M-IMP-2", "Қайталанған дербес шот", "Мекенжай 2", ""],
        ["0222222", "M-IMP-1", "Қайталанған есептегіш", "Мекенжай 3", ""],
        ["bad", "M-IMP-3", "Жарамсыз шот", "Мекенжай 4", ""],
        ["0777777", "M-IMP-4", "Бұрыннан бар шот", "Мекенжай 5", ""],
        ["", "", "", "", ""],
    ])
    response = await logged_in.post("/api/subscribers/import",
        files={"file": ("import.xlsx", content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["created"] == 1
    assert len(body["errors"]) == 3  # duplicate account, duplicate meter within file, invalid account format
    assert len(body["skipped"]) == 1  # already registered account
    stored = await db.scalar(select(SubscriberAccount).where(SubscriberAccount.account_number == "0111111"))
    assert stored is not None and stored.full_name == "Бірінші Т."
    total = await db.scalar(select(func.count()).select_from(SubscriberAccount))
    assert total == 3  # conftest's default account + the pre-registered 0777777 + the one new valid row


async def test_import_csv_windows1251_encoding(logged_in):
    # Windows-1251 has no Kazakh-specific letters, so a real cp1251 export uses Russian labels/text.
    text = "Лицевой счет,Номер счетчика,ФИО,Адрес,Телефон\n0333333,M-CSV-1,Тестов Т.,Тестовая улица,\n"
    content = text.encode("cp1251")
    response = await logged_in.post("/api/subscribers/import", files={"file": ("x.csv", content, "text/csv")})
    assert response.status_code == 200, response.text
    assert response.json()["created"] == 1


async def test_manual_binding_concurrency_and_readings(logged_in, db):
    row = await register(logged_in, phone=None)
    binding = {"telegram_user_id": 777, "version": row["version"], "reason": "Identity checked in person"}
    assert (await logged_in.post(f"/api/subscribers/{row['id']}/binding", json=binding)).status_code == 200
    assert (await logged_in.post(f"/api/subscribers/{row['id']}/binding", json=binding)).status_code == 409
    period = local_today().replace(day=1).isoformat()
    photo_id = await upload_reading_photo(db, 777)
    payload = {"account_id": row["id"], "telegram_user_id": 777, "idempotency_key": str(uuid.uuid4()),
               "period": period, "photo_id": photo_id}
    wrong = await logged_in.post("/api/internal/subscribers/readings", headers=bot_headers(),
                                json={**payload, "telegram_user_id": 888})
    assert wrong.status_code == 403
    first = await logged_in.post("/api/internal/subscribers/readings", headers=bot_headers(), json=payload)
    assert first.status_code == 201, first.text
    again = await logged_in.post("/api/internal/subscribers/readings", headers=bot_headers(), json=payload)
    assert again.json()["id"] == first.json()["id"]
    other_photo_id = await upload_reading_photo(db, 777)
    changed = await logged_in.post("/api/internal/subscribers/readings", headers=bot_headers(),
                                   json={**payload, "photo_id": other_photo_id})
    assert changed.status_code == 409
    review = await logged_in.post(f"/api/subscribers/readings/{first.json()['id']}/review",
                                   json={"status": "ACCEPTED", "note": "Initial baseline checked", "value": "1234.125"})
    assert review.status_code == 200 and review.json()["consumption"] is None
    stored = await db.get(SubscriberAccount, row["id"])
    assert stored.last_reading == Decimal("1234.125")
    repeat_review = await logged_in.post(f"/api/subscribers/readings/{first.json()['id']}/review",
                                         json={"status": "ACCEPTED", "note": "Repeat", "value": "1234.125"})
    assert repeat_review.status_code == 409
    assert (await logged_in.get("/api/internal/subscribers/888/readings", headers=bot_headers())).json() == {"items": []}


async def test_consumption_from_previous_accepted_value(logged_in, db):
    account = await db.scalar(select(SubscriberAccount).where(SubscriberAccount.account_number == "12345678"))
    account.last_reading = Decimal("100.125")
    period = local_today().replace(day=1)
    account.last_period = period.replace(year=period.year - 1)
    await db.commit()
    photo_id = await upload_reading_photo(db, 123456)
    payload = {"account_id": account.id, "telegram_user_id": 123456, "idempotency_key": str(uuid.uuid4()),
               "period": period.isoformat(), "photo_id": photo_id}
    reading = await logged_in.post("/api/internal/subscribers/readings", headers=bot_headers(), json=payload)
    assert reading.status_code == 201
    assert reading.json()["value"] is None
    rejected_low = await logged_in.post(f"/api/subscribers/readings/{reading.json()['id']}/review",
                                        json={"status": "ACCEPTED", "note": "Too low", "value": "50.000"})
    assert rejected_low.status_code == 422
    accepted = await logged_in.post(f"/api/subscribers/readings/{reading.json()['id']}/review",
                                    json={"status": "ACCEPTED", "note": "Checked", "value": "110.250"})
    assert Decimal(str(accepted.json()["consumption"])) == Decimal("10.125")
    assert (await db.get(MeterReading, reading.json()["id"])).status == "ACCEPTED"


async def test_application_api_cannot_bypass_binding(logged_in):
    from test_backend import payload
    response = await logged_in.post("/api/internal/applications", headers=bot_headers(),
                                    json=payload(personal_account="99999999"))
    assert response.status_code == 403


@pytest.mark.parametrize("value", ["NaN", "Infinity", "-1", "1.0001"])
async def test_invalid_review_values_rejected(logged_in, db, value):
    photo_id = await upload_reading_photo(db, 123456)
    submitted = await logged_in.post("/api/internal/subscribers/readings", headers=bot_headers(), json={
        "account_id": 1, "telegram_user_id": 123456, "idempotency_key": str(uuid.uuid4()),
        "period": local_today().replace(day=1).isoformat(), "photo_id": photo_id})
    assert submitted.status_code == 201
    response = await logged_in.post(f"/api/subscribers/readings/{submitted.json()['id']}/review",
                                    json={"status": "ACCEPTED", "note": "Checked", "value": value})
    assert response.status_code == 422


async def test_submit_reading_rejects_photo_that_is_not_a_meter(logged_in, db, monkeypatch, tmp_path):
    from unittest.mock import AsyncMock
    monkeypatch.setattr(get_settings(), "upload_dir", tmp_path)
    monkeypatch.setattr("app.subscribers.looks_like_gas_meter", AsyncMock(return_value=False))
    photo_id = await upload_reading_photo(db, 123456)
    photo = await db.get(ApplicationFile, photo_id)
    (tmp_path / photo.storage_url).write_bytes(b"not a meter")
    response = await logged_in.post("/api/internal/subscribers/readings", headers=bot_headers(), json={
        "account_id": 1, "telegram_user_id": 123456, "idempotency_key": str(uuid.uuid4()),
        "period": local_today().replace(day=1).isoformat(), "photo_id": photo_id})
    assert response.status_code == 422


async def test_submit_reading_allowed_when_photo_confirmed_as_meter(logged_in, db, monkeypatch, tmp_path):
    from unittest.mock import AsyncMock
    monkeypatch.setattr(get_settings(), "upload_dir", tmp_path)
    monkeypatch.setattr("app.subscribers.looks_like_gas_meter", AsyncMock(return_value=True))
    photo_id = await upload_reading_photo(db, 123456)
    photo = await db.get(ApplicationFile, photo_id)
    (tmp_path / photo.storage_url).write_bytes(b"meter photo")
    response = await logged_in.post("/api/internal/subscribers/readings", headers=bot_headers(), json={
        "account_id": 1, "telegram_user_id": 123456, "idempotency_key": str(uuid.uuid4()),
        "period": local_today().replace(day=1).isoformat(), "photo_id": photo_id})
    assert response.status_code == 201, response.text


async def test_reject_photo_reused_from_an_earlier_period(logged_in, db):
    shared_hash = "ab" * 32
    period = local_today().replace(day=1)
    old_period = period.replace(year=period.year - 1) if period.month == 1 else period.replace(month=period.month - 1)
    old_photo_id = await upload_reading_photo(db, 123456, content_hash=shared_hash)
    db.add(MeterReading(account_id=1, telegram_user_id=123456, idempotency_key=str(uuid.uuid4()),
                        period=old_period, photo_id=old_photo_id, status="ACCEPTED", value=Decimal("5.000")))
    await db.commit()
    new_photo_id = await upload_reading_photo(db, 123456, content_hash=shared_hash)

    response = await logged_in.post("/api/internal/subscribers/readings", headers=bot_headers(), json={
        "account_id": 1, "telegram_user_id": 123456, "idempotency_key": str(uuid.uuid4()),
        "period": period.isoformat(), "photo_id": new_photo_id})
    assert response.status_code == 422
    assert "ертерек жіберілген" in response.text


async def test_reading_rejects_foreign_or_wrong_type_photo(logged_in, db):
    foreign_photo = await upload_reading_photo(db, 999999)
    wrong_type = ApplicationFile(owner_telegram_id=123456, telegram_file_id="tg-" + str(uuid.uuid4()),
                                 file_type=FileType.METER_PHOTO, storage_url=str(uuid.uuid4()))
    db.add(wrong_type)
    await db.commit()
    base = {"account_id": 1, "telegram_user_id": 123456,
            "period": local_today().replace(day=1).isoformat()}
    foreign = await logged_in.post("/api/internal/subscribers/readings", headers=bot_headers(),
                                   json={**base, "idempotency_key": str(uuid.uuid4()), "photo_id": foreign_photo})
    assert foreign.status_code == 422
    mismatched = await logged_in.post("/api/internal/subscribers/readings", headers=bot_headers(),
                                      json={**base, "idempotency_key": str(uuid.uuid4()), "photo_id": str(wrong_type.id)})
    assert mismatched.status_code == 422


async def test_review_requires_value_when_accepting(logged_in, db):
    photo_id = await upload_reading_photo(db, 123456)
    submitted = await logged_in.post("/api/internal/subscribers/readings", headers=bot_headers(), json={
        "account_id": 1, "telegram_user_id": 123456, "idempotency_key": str(uuid.uuid4()),
        "period": local_today().replace(day=1).isoformat(), "photo_id": photo_id})
    assert submitted.status_code == 201
    missing_value = await logged_in.post(f"/api/subscribers/readings/{submitted.json()['id']}/review",
                                         json={"status": "ACCEPTED", "note": "Checked"})
    assert missing_value.status_code == 422
    rejected = await logged_in.post(f"/api/subscribers/readings/{submitted.json()['id']}/review",
                                    json={"status": "REJECTED", "note": "Blurry photo"})
    assert rejected.status_code == 200 and rejected.json()["value"] is None


async def test_postgres_concurrent_contact_only_one_owner(db, monkeypatch):
    import asyncio
    from unittest.mock import AsyncMock
    from sqlalchemy.ext.asyncio import async_sessionmaker
    from app.subscribers import ContactProof, verify_contact
    if db.bind.dialect.name != "postgresql":
        pytest.skip("Requires PostgreSQL row locks")
    monkeypatch.setattr("app.subscribers.rate_limit", AsyncMock())
    row = SubscriberAccount(account_number="00112233", meter_number="CONTACT-RACE", full_name="Synthetic resident",
                            address="Test address", phone="+77010000001")
    db.add(row)
    await db.commit()
    factory = async_sessionmaker(db.bind, expire_on_commit=False)

    async def bind(user_id):
        async with factory() as session:
            return await verify_contact(ContactProof(telegram_user_id=user_id, contact_user_id=user_id,
                identifier="00112233", phone="+77010000001"), session)

    results = await asyncio.gather(bind(111), bind(222))
    assert sorted(result["verified"] for result in results) == [False, True]


async def test_postgres_concurrent_readings_one_report(db, monkeypatch):
    import asyncio
    from unittest.mock import AsyncMock
    from sqlalchemy.ext.asyncio import async_sessionmaker
    from app.subscribers import ReadingInput, submit_reading
    if db.bind.dialect.name != "postgresql":
        pytest.skip("Requires PostgreSQL row locks")
    monkeypatch.setattr("app.subscribers.rate_limit", AsyncMock())
    account = await db.scalar(select(SubscriberAccount).where(SubscriberAccount.account_number == "12345678"))
    photo_id = await upload_reading_photo(db, 123456)
    await db.commit()
    factory = async_sessionmaker(db.bind, expire_on_commit=False)
    payload = ReadingInput(account_id=account.id, telegram_user_id=123456, idempotency_key=uuid.uuid4(),
                           period=local_today().replace(day=1), photo_id=photo_id)

    async def submit():
        async with factory() as session:
            return await submit_reading(payload, session)

    results = await asyncio.gather(*[submit() for _ in range(5)])
    assert len({result["id"] for result in results}) == 1


async def test_reminder_sent_once_per_month_and_skipped_for_pending(db):
    from app.models import Outbox
    from app.notifications import send_reminders

    period = local_today().replace(day=1)
    due = SubscriberAccount(account_number="0400001", meter_number="M-REM-1", full_name="Due resident",
                            address="Test address", telegram_user_id=900001, active=True)
    already_reported = SubscriberAccount(account_number="0400002", meter_number="M-REM-2", full_name="Reported resident",
                                         address="Test address", telegram_user_id=900002, active=True, last_period=period)
    unbound = SubscriberAccount(account_number="0400003", meter_number="M-REM-3", full_name="Unbound resident",
                                address="Test address", telegram_user_id=None, active=True)
    inactive = SubscriberAccount(account_number="0400004", meter_number="M-REM-4", full_name="Inactive resident",
                                 address="Test address", telegram_user_id=900004, active=False)
    db.add_all([due, already_reported, unbound, inactive])
    await db.commit()
    pending_photo = await upload_reading_photo(db, 900005)
    awaiting_review = SubscriberAccount(account_number="0400005", meter_number="M-REM-5", full_name="Awaiting resident",
                                        address="Test address", telegram_user_id=900005, active=True)
    db.add(awaiting_review)
    await db.commit()
    db.add(MeterReading(account_id=awaiting_review.id, telegram_user_id=900005, idempotency_key=str(uuid.uuid4()),
                        period=period, photo_id=pending_photo, status="PENDING"))
    await db.commit()

    await send_reminders(db)

    def sent_to(user_id):
        return db.scalars(select(Outbox).where(Outbox.telegram_user_id == user_id))

    due_messages = (await sent_to(900001)).all()
    assert len(due_messages) == 1 and "0400001" in due_messages[0].text
    assert (await sent_to(900002)).all() == []  # already reported this period
    assert (await sent_to(900004)).all() == []  # inactive
    assert (await sent_to(900005)).all() == []  # has a pending report awaiting review

    await db.refresh(due)
    assert due.last_reminder_period == period
    await db.refresh(awaiting_review)
    assert awaiting_review.last_reminder_period == period  # marked done without nagging

    # A second run this same period must not send a duplicate.
    await send_reminders(db)
    assert len((await sent_to(900001)).all()) == 1


async def test_registry_stats_counts_by_period(logged_in, db):
    period = local_today().replace(day=1)
    reported = SubscriberAccount(account_number="0500001", meter_number="M-STAT-1", full_name="Reported",
                                 address="Addr", telegram_user_id=910001, active=True, last_period=period)
    awaiting = SubscriberAccount(account_number="0500002", meter_number="M-STAT-2", full_name="Awaiting",
                                 address="Addr", telegram_user_id=910002, active=True)
    missing = SubscriberAccount(account_number="0500003", meter_number="M-STAT-3", full_name="Missing",
                                address="Addr", telegram_user_id=910003, active=True)
    unbound = SubscriberAccount(account_number="0500004", meter_number="M-STAT-4", full_name="Unbound",
                                address="Addr", telegram_user_id=None, active=True)
    db.add_all([reported, awaiting, missing, unbound])
    await db.commit()
    photo_id = await upload_reading_photo(db, 910002)
    db.add(MeterReading(account_id=awaiting.id, telegram_user_id=910002, idempotency_key=str(uuid.uuid4()),
                        period=period, photo_id=photo_id, status="PENDING"))
    await db.commit()

    response = await logged_in.get("/api/subscribers/stats")
    assert response.status_code == 200, response.text
    body = response.json()
    # conftest's default account (12345678, telegram_user_id=123456) is bound and unreported too,
    # so it counts alongside "missing" here: total 4 = default + reported + awaiting + missing.
    assert body == {"period": period.isoformat(), "total": 4, "accepted": 1, "pending": 1, "missing": 2}


async def test_export_escapes_formulas_in_both_formats(logged_in):
    await register(logged_in, account_number="0600001", meter_number="M-EXP-1", full_name="=cmd|'/c calc'!A1")
    csv_response = await logged_in.get("/api/subscribers/export?format=csv")
    assert csv_response.status_code == 200
    assert csv_response.headers["content-type"].startswith("text/csv")
    assert "'=cmd" in csv_response.text
    xlsx_response = await logged_in.get("/api/subscribers/export?format=xlsx")
    assert xlsx_response.status_code == 200
    assert xlsx_response.headers["content-type"].startswith("application/vnd.openxmlformats")


async def test_export_requires_admin(client):
    denied = await client.get("/api/subscribers/export")
    assert denied.status_code == 401


async def test_account_profile_aggregates_readings_applications_and_history(logged_in, db):
    from test_backend import payload

    account = await db.scalar(select(SubscriberAccount).where(SubscriberAccount.account_number == "12345678"))
    # Send back the exact current values: a no-op update still logs "subscriber_update"
    # without touching full_name/address/phone/active, which would otherwise unbind Telegram.
    updated = await logged_in.put(f"/api/subscribers/{account.id}", json={
        "account_number": "12345678", "meter_number": "TEST-METER-1", "full_name": "Test resident",
        "address": "Test address", "phone": None, "version": account.version, "active": True})
    assert updated.status_code == 200, updated.text

    photo_id = await upload_reading_photo(db, 123456)
    reading = await logged_in.post("/api/internal/subscribers/readings", headers=bot_headers(), json={
        "account_id": account.id, "telegram_user_id": 123456, "idempotency_key": str(uuid.uuid4()),
        "period": local_today().replace(day=1).isoformat(), "photo_id": photo_id})
    assert reading.status_code == 201, reading.text

    created = await logged_in.post("/api/internal/applications", headers=bot_headers(), json=payload())
    assert created.status_code == 201, created.text

    profile = await logged_in.get(f"/api/subscribers/{account.id}/profile")
    assert profile.status_code == 200, profile.text
    body = profile.json()
    assert body["account"]["account_number"] == "12345678"
    assert len(body["readings"]) == 1
    assert len(body["applications"]) == 1
    assert body["applications"][0]["personal_account"] == "12345678"
    assert any(h["action"] == "Деректер өзгертілді" for h in body["history"])


async def test_account_profile_not_found(logged_in):
    response = await logged_in.get("/api/subscribers/999999/profile")
    assert response.status_code == 404
